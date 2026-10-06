"""Professional Excel Exporter for New York Sales Tax Report (NY_Sales_Tax_Report.xlsx)."""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from tax_tools.sales_tax_models import SalesTaxReport, TransactionTreatment


# Typography & Palette Constants
FONT_FAMILY = "Segoe UI"

# Colors
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")       # Dark Navy
SUBHEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")    # Soft Blue
CARD_ACCENT_FILL = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")  # Medium Navy
TOTAL_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")        # Accent Blue
SUCCESS_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")      # Soft Green
TAXABLE_FILL = PatternFill(start_color="EBF1F5", end_color="EBF1F5", fill_type="solid")      # Very Soft Blue
REFUND_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")       # Soft Peach
EXCLUDED_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")     # Subtle Gray
KPI_BG_FILL = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")       # Clean Off-White

TITLE_FONT = Font(name=FONT_FAMILY, size=16, bold=True, color="1F4E79")
SECTION_FONT = Font(name=FONT_FAMILY, size=12, bold=True, color="1F4E79")
HEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
SUBHEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="1F4E79")
REGULAR_FONT = Font(name=FONT_FAMILY, size=9, bold=False, color="000000")
BOLD_FONT = Font(name=FONT_FAMILY, size=9, bold=True, color="000000")
KPI_TITLE_FONT = Font(name=FONT_FAMILY, size=9, bold=False, color="595959")
KPI_VALUE_FONT = Font(name=FONT_FAMILY, size=14, bold=True, color="1F4E79")
KPI_TAX_VALUE_FONT = Font(name=FONT_FAMILY, size=16, bold=True, color="1E4620")

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
KPI_BOX_BORDER = Border(
    left=Side(style="medium", color="BDD7EE"),
    right=Side(style="medium", color="BDD7EE"),
    top=Side(style="medium", color="BDD7EE"),
    bottom=Side(style="medium", color="BDD7EE"),
)
KPI_TAX_BOX_BORDER = Border(
    left=Side(style="medium", color="A9D18E"),
    right=Side(style="medium", color="A9D18E"),
    top=Side(style="medium", color="A9D18E"),
    bottom=Side(style="medium", color="A9D18E"),
)
DOUBLE_BOTTOM_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="double", color="000000"),
)

CURRENCY_FORMAT = "$#,##0.00;($#,##0.00);\"-\""
PERCENT_FORMAT = "0.000%"


class SalesTaxExcelExporter:
    """Exports SalesTaxReport into a structured, audit-ready 3-tab Excel workbook."""

    def export(self, report: SalesTaxReport, output_path: str = "NY_Sales_Tax_Report.xlsx") -> str:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default blank sheet

        ws1 = wb.create_sheet(title="Deduction Summary")
        self._build_deduction_summary_tab(ws1, report)

        ws2 = wb.create_sheet(title="Itemized Ledger")
        self._build_itemized_ledger_tab(ws2, report)

        ws3 = wb.create_sheet(title="Category Audit Trail")
        self._build_category_audit_tab(ws3, report)

        wb.save(output_path)
        return output_path

    def _build_deduction_summary_tab(self, ws, report: SalesTaxReport):
        ws.views.sheetView[0].showGridLines = True

        # Header Block
        ws["A1"] = "NEW YORK SALES TAX ITEMIZATION SUMMARY"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = "IRS Schedule A / NY State Sales & Use Tax Itemized Deduction Report"
        ws["A2"].font = Font(name=FONT_FAMILY, size=11, italic=True, color="595959")

        ws["A4"] = "Tax Jurisdiction:"
        ws["B4"] = report.county
        ws["A5"] = "Combined Statutory Rate:"
        ws["B5"] = report.tax_rate
        ws["B5"].number_format = PERCENT_FORMAT
        ws["A6"] = "Total Statements / Transactions:"
        ws["B6"] = report.total_transactions

        for row in range(4, 7):
            ws[f"A{row}"].font = BOLD_FONT
            ws[f"B{row}"].font = REGULAR_FONT

        # -------------------------------------------------------------
        # Row 8-12: KPI Summary Cards (Grid Layout)
        # -------------------------------------------------------------
        # Card 1: Gross Spend Analyzed
        self._render_kpi_card(
            ws, start_col=1, start_row=8, width=2,
            title="TOTAL SPEND ANALYZED",
            value=report.total_gross_spend,
            is_currency=True,
        )
        # Card 2: Non-Spend Excluded
        self._render_kpi_card(
            ws, start_col=3, start_row=8, width=2,
            title="NON-SPEND EXCLUDED",
            value=report.total_excluded,
            is_currency=True,
        )
        # Card 3: Tax-Exempt Purchases
        self._render_kpi_card(
            ws, start_col=5, start_row=8, width=2,
            title="TAX-EXEMPT PURCHASES",
            value=report.total_exempt,
            is_currency=True,
        )

        # Card 4: Out-of-State / Foreign
        self._render_kpi_card(
            ws, start_col=1, start_row=12, width=2,
            title="OUT-OF-STATE EXCLUDED",
            value=report.total_out_of_state,
            is_currency=True,
        )
        # Card 5: Taxable Spend (Net)
        self._render_kpi_card(
            ws, start_col=3, start_row=12, width=2,
            title="NET TAXABLE SPEND",
            value=report.total_taxable,
            is_currency=True,
        )
        # Card 6: Deductible NY Sales Tax Paid (HERO CARD)
        self._render_kpi_card(
            ws, start_col=5, start_row=12, width=2,
            title="DEDUCTIBLE NY SALES TAX PAID",
            value=report.total_sales_tax_paid,
            is_currency=True,
            is_hero=True,
        )

        # -------------------------------------------------------------
        # Row 16+: Itemized Deduction Reconciliation Table
        # -------------------------------------------------------------
        start_r = 16
        ws.cell(row=start_r, column=1, value="SCHEDULE A SALES TAX DEDUCTION RECONCILIATION").font = SECTION_FONT

        headers = ["Reconciliation Line Item", "Statutory Classification", "Amount ($)"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=start_r + 1, column=col_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left")

        rows_data = [
            ("Gross Spending Volume Incurred", "All goods, services, and transactions", report.total_gross_spend),
            ("Less: Sales-Tax-Exempt Spend", "Unprepared food, medicine, transit, utilities", -report.total_exempt),
            ("Less: Out-of-State / Foreign Spend", "Airfare, lodging outside NY, foreign currency", -report.total_out_of_state),
            ("Net Qualifying Taxable Spend (Inclusive of Tax)", "Subject to reverse sales tax calculation", report.total_taxable),
            ("Reverse-Calculated Pre-Tax Purchase Base", "Amount / (1 + Rate)", report.total_pre_tax),
            ("TOTAL DEDUCTIBLE NY SALES TAX PAID", "Amount * (Rate / (1 + Rate))", report.total_sales_tax_paid),
        ]

        curr_r = start_r + 2
        for idx, (label, desc, amt) in enumerate(rows_data):
            c1 = ws.cell(row=curr_r, column=1, value=label)
            c2 = ws.cell(row=curr_r, column=2, value=desc)
            c3 = ws.cell(row=curr_r, column=3, value=amt)

            c1.border = THIN_BORDER
            c2.border = THIN_BORDER
            c3.border = THIN_BORDER

            c3.number_format = CURRENCY_FORMAT
            c3.alignment = Alignment(horizontal="right")

            if idx == len(rows_data) - 1:  # Grand total row
                c1.fill = SUCCESS_FILL
                c2.fill = SUCCESS_FILL
                c3.fill = SUCCESS_FILL
                c1.font = BOLD_FONT
                c2.font = BOLD_FONT
                c3.font = KPI_TAX_VALUE_FONT
                c1.border = DOUBLE_BOTTOM_BORDER
                c2.border = DOUBLE_BOTTOM_BORDER
                c3.border = DOUBLE_BOTTOM_BORDER
            elif idx == 3:  # Net taxable spend
                c1.fill = SUBHEADER_FILL
                c2.fill = SUBHEADER_FILL
                c3.fill = SUBHEADER_FILL
                c1.font = BOLD_FONT
                c2.font = REGULAR_FONT
                c3.font = BOLD_FONT
            else:
                c1.font = REGULAR_FONT
                c2.font = REGULAR_FONT
                c3.font = REGULAR_FONT

            curr_r += 1

        self._autofit_columns(ws)

    def _render_kpi_card(self, ws, start_col: int, start_row: int, width: int, title: str, value: float, is_currency: bool = True, is_hero: bool = False):
        end_col = start_col + width - 1

        ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=end_col)
        t_cell = ws.cell(row=start_row, column=start_col, value=title)
        t_cell.font = KPI_TITLE_FONT
        t_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells(start_row=start_row + 1, start_column=start_col, end_row=start_row + 2, end_column=end_col)
        v_cell = ws.cell(row=start_row + 1, column=start_col, value=value)
        v_cell.font = KPI_TAX_VALUE_FONT if is_hero else KPI_VALUE_FONT
        v_cell.alignment = Alignment(horizontal="center", vertical="center")
        if is_currency:
            v_cell.number_format = CURRENCY_FORMAT

        box_fill = SUCCESS_FILL if is_hero else KPI_BG_FILL
        border_style = KPI_TAX_BOX_BORDER if is_hero else KPI_BOX_BORDER

        for r in range(start_row, start_row + 3):
            for c in range(start_col, end_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = box_fill
                cell.border = border_style

    def _build_itemized_ledger_tab(self, ws, report: SalesTaxReport):
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        headers = [
            "Date",
            "Bank",
            "Account",
            "Description",
            "Settled Amount ($)",
            "Status",
            "Category",
            "Pre-Tax Base ($)",
            "Sales Tax ($)",
            "Tax Rate",
            "Audit Rationale / Statutory Authority",
        ]

        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = THIN_BORDER

        row_idx = 2
        for t in report.transactions:
            c1 = ws.cell(row=row_idx, column=1, value=t.date)
            c2 = ws.cell(row=row_idx, column=2, value=t.source_bank)
            c3 = ws.cell(row=row_idx, column=3, value=t.source_account)
            c4 = ws.cell(row=row_idx, column=4, value=t.description)
            c5 = ws.cell(row=row_idx, column=5, value=t.amount)
            c6 = ws.cell(row=row_idx, column=6, value=t.treatment.value)
            c7 = ws.cell(row=row_idx, column=7, value=t.category)
            c8 = ws.cell(row=row_idx, column=8, value=t.pre_tax_amount if t.is_taxable else 0.0)
            c9 = ws.cell(row=row_idx, column=9, value=t.sales_tax if t.is_taxable else 0.0)
            c10 = ws.cell(row=row_idx, column=10, value=t.tax_rate if t.is_taxable else 0.0)
            c11 = ws.cell(row=row_idx, column=11, value=t.classification_reason)

            # Alignments
            c1.alignment = Alignment(horizontal="center")
            c5.alignment = Alignment(horizontal="right")
            c6.alignment = Alignment(horizontal="center")
            c8.alignment = Alignment(horizontal="right")
            c9.alignment = Alignment(horizontal="right")
            c10.alignment = Alignment(horizontal="right")

            # Formats
            c5.number_format = CURRENCY_FORMAT
            c8.number_format = CURRENCY_FORMAT
            c9.number_format = CURRENCY_FORMAT
            c10.number_format = PERCENT_FORMAT

            # Row fill styling based on status
            fill = None
            if t.treatment == TransactionTreatment.TAXABLE:
                fill = TAXABLE_FILL
            elif t.treatment == TransactionTreatment.REFUND:
                fill = REFUND_FILL
            elif t.is_excluded:
                fill = EXCLUDED_FILL

            for cell in (c1, c2, c3, c4, c5, c6, c7, c8, c9, c10, c11):
                cell.font = REGULAR_FONT
                cell.border = THIN_BORDER
                if fill:
                    cell.fill = fill

            row_idx += 1

        self._autofit_columns(ws, max_width_override={"Audit Rationale / Statutory Authority": 65})

    def _build_category_audit_tab(self, ws, report: SalesTaxReport):
        ws.views.sheetView[0].showGridLines = True

        ws["A1"] = "CATEGORY & INSTITUTION AUDIT TRAIL"
        ws["A1"].font = TITLE_FONT

        # Table 1: Category Breakdown
        ws["A3"] = "SPEND BREAKDOWN BY CATEGORY"
        ws["A3"].font = SECTION_FONT

        cat_headers = [
            "Category",
            "Tax Treatment",
            "Count",
            "Gross Spend ($)",
            "Taxable Spend ($)",
            "Sales Tax Paid ($)",
            "% of Total Tax",
        ]
        for col_idx, h in enumerate(cat_headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = THIN_BORDER

        curr_r = 5
        for cs in report.category_summaries:
            c1 = ws.cell(row=curr_r, column=1, value=cs.category)
            c2 = ws.cell(row=curr_r, column=2, value=cs.treatment)
            c3 = ws.cell(row=curr_r, column=3, value=cs.transaction_count)
            c4 = ws.cell(row=curr_r, column=4, value=cs.gross_amount)
            c5 = ws.cell(row=curr_r, column=5, value=cs.taxable_amount)
            c6 = ws.cell(row=curr_r, column=6, value=cs.sales_tax_paid)
            c7 = ws.cell(row=curr_r, column=7, value=cs.pct_of_total_tax / 100.0 if cs.pct_of_total_tax else 0.0)

            c2.alignment = Alignment(horizontal="center")
            c3.alignment = Alignment(horizontal="right")
            c4.alignment = Alignment(horizontal="right")
            c5.alignment = Alignment(horizontal="right")
            c6.alignment = Alignment(horizontal="right")
            c7.alignment = Alignment(horizontal="right")

            c4.number_format = CURRENCY_FORMAT
            c5.number_format = CURRENCY_FORMAT
            c6.number_format = CURRENCY_FORMAT
            c7.number_format = "0.0%"

            for c in (c1, c2, c3, c4, c5, c6, c7):
                c.font = REGULAR_FONT
                c.border = THIN_BORDER

            curr_r += 1

        # Total Row for Category Table
        tot_c1 = ws.cell(row=curr_r, column=1, value="Total Spend & Deductions")
        tot_c2 = ws.cell(row=curr_r, column=2, value="")
        tot_c3 = ws.cell(row=curr_r, column=3, value=sum(cs.transaction_count for cs in report.category_summaries))
        tot_c4 = ws.cell(row=curr_r, column=4, value=round(sum(cs.gross_amount for cs in report.category_summaries), 2))
        tot_c5 = ws.cell(row=curr_r, column=5, value=report.total_taxable)
        tot_c6 = ws.cell(row=curr_r, column=6, value=report.total_sales_tax_paid)
        tot_c7 = ws.cell(row=curr_r, column=7, value=1.0 if report.total_sales_tax_paid > 0 else 0.0)

        for c in (tot_c1, tot_c2, tot_c3, tot_c4, tot_c5, tot_c6, tot_c7):
            c.font = BOLD_FONT
            c.fill = TOTAL_FILL
            c.border = DOUBLE_BOTTOM_BORDER

        tot_c4.number_format = CURRENCY_FORMAT
        tot_c5.number_format = CURRENCY_FORMAT
        tot_c6.number_format = CURRENCY_FORMAT
        tot_c7.number_format = "0.0%"

        # Table 2: Card & Account Breakdown
        curr_r += 3
        ws.cell(row=curr_r, column=1, value="SPEND BREAKDOWN BY INSTITUTION & ACCOUNT").font = SECTION_FONT
        curr_r += 1

        card_headers = [
            "Source Bank",
            "Card / Account",
            "Count",
            "Gross Spend ($)",
            "Taxable Spend ($)",
            "Sales Tax Paid ($)",
        ]
        for col_idx, h in enumerate(card_headers, 1):
            cell = ws.cell(row=curr_r, column=col_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = CARD_ACCENT_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = THIN_BORDER

        curr_r += 1
        for cds in report.card_summaries:
            c1 = ws.cell(row=curr_r, column=1, value=cds.source_bank)
            c2 = ws.cell(row=curr_r, column=2, value=cds.source_account)
            c3 = ws.cell(row=curr_r, column=3, value=cds.transaction_count)
            c4 = ws.cell(row=curr_r, column=4, value=cds.gross_amount)
            c5 = ws.cell(row=curr_r, column=5, value=cds.taxable_amount)
            c6 = ws.cell(row=curr_r, column=6, value=cds.sales_tax_paid)

            c3.alignment = Alignment(horizontal="right")
            c4.alignment = Alignment(horizontal="right")
            c5.alignment = Alignment(horizontal="right")
            c6.alignment = Alignment(horizontal="right")

            c4.number_format = CURRENCY_FORMAT
            c5.number_format = CURRENCY_FORMAT
            c6.number_format = CURRENCY_FORMAT

            for c in (c1, c2, c3, c4, c5, c6):
                c.font = REGULAR_FONT
                c.border = THIN_BORDER

            curr_r += 1

        self._autofit_columns(ws)

    def _autofit_columns(self, ws, max_width_override: dict = None):
        """Auto-fit column widths with padding and maximum thresholds."""
        max_overrides = max_width_override or {}
        for col in ws.columns:
            header_val = str(col[0].value or "")
            max_len = 0
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)

            col_letter = get_column_letter(col[0].column)
            # Default width bounded between 12 and 50 unless overridden
            capped_width = min(max(max_len + 4, 12), max_overrides.get(header_val, 50))
            ws.column_dimensions[col_letter].width = capped_width
