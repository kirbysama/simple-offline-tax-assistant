"""Excel Exporter for Form 8949 Tax Summary and Brokerage Breakdown."""

import os
from typing import Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from tax_tools.models import AggregationReport, Form8949Summary


# Palette & Styling Constants
FONT_FAMILY = "Segoe UI"
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
SUCCESS_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

# Wash Sale Highlighting Constants (Soft Red fill, bold dark red font)
WASH_SALE_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
WASH_SALE_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="9C0006")

TITLE_FONT = Font(name=FONT_FAMILY, size=16, bold=True, color="1F4E79")
HEADER_FONT = Font(name=FONT_FAMILY, size=11, bold=True, color="FFFFFF")
SUBHEADER_FONT = Font(name=FONT_FAMILY, size=11, bold=True, color="1F4E79")
REGULAR_FONT = Font(name=FONT_FAMILY, size=10, bold=False, color="000000")
BOLD_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="000000")
SUCCESS_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="375623")

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
DOUBLE_BOTTOM_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="double", color="000000"),
)

CURRENCY_FORMAT = "$#,##0.00;($#,##0.00);\"-\""


def export_1099_excel(report: AggregationReport, output_path: str = "1099_Tax_Summary.xlsx") -> str:
    """
    Exports Form 8949 Summary, Brokerage Breakdown, and Interest & Dividend Schedule to a 3-tab Excel workbook.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Tab 1: Form 8949 Summary
    ws_summary = wb.create_sheet(title="Form 8949 Summary")
    _build_form_8949_tab(ws_summary, report.aggregate_8949)

    # Tab 2: Brokerage Breakdown
    ws_breakdown = wb.create_sheet(title="Brokerage Breakdown")
    _build_breakdown_tab(ws_breakdown, report)

    # Tab 3: Interest & Dividend Schedule
    ws_int_div = wb.create_sheet(title="Interest & Dividend Schedule")
    _build_int_div_tab(ws_int_div, report)

    # Save workbook
    wb.save(output_path)
    return output_path


def _build_form_8949_tab(ws, summary: Form8949Summary):
    """Builds the IRS Form 8949 / Schedule D Summary tab."""
    ws.views.sheetView[0].showGridLines = True

    # Title block
    ws["A1"] = "Consolidated IRS Form 8949 / Schedule D Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Deterministic extraction and arithmetic reconciliation across all brokerage accounts"
    ws["A2"].font = Font(name=FONT_FAMILY, size=10, italic=True, color="595959")

    headers = [
        "Box",
        "Form 8949 Category Description",
        "Proceeds (1d)",
        "Cost Basis (1e)",
        "Adjustments (1g/1f)",
        "Net Gain / (Loss)",
        "Parity Status",
    ]

    row_num = 4
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=row_num, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center" if col_idx in (1, 7) else "left", vertical="center")
        cell.border = THIN_BORDER
    ws.row_dimensions[row_num].height = 24

    # Part I Header
    row_num += 1
    cell = ws.cell(row=row_num, column=1, value="Part I: Short-Term Capital Gains & Losses (Assets held one year or less)")
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=len(headers))
    cell.font = SUBHEADER_FONT
    cell.fill = SUBHEADER_FILL
    ws.row_dimensions[row_num].height = 20

    # Part I Rows: Box A, B, C
    part_i_boxes = [summary.box_a, summary.box_b, summary.box_c]
    for box in part_i_boxes:
        row_num += 1
        _write_box_row(ws, row_num, box)

    # Part I Subtotal
    row_num += 1
    _write_box_row(ws, row_num, summary.short_term_total, is_subtotal=True)

    # Part II Header
    row_num += 2
    cell = ws.cell(row=row_num, column=1, value="Part II: Long-Term Capital Gains & Losses (Assets held more than one year)")
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=len(headers))
    cell.font = SUBHEADER_FONT
    cell.fill = SUBHEADER_FILL
    ws.row_dimensions[row_num].height = 20

    # Part II Rows: Box D, E, F
    part_ii_boxes = [summary.box_d, summary.box_e, summary.box_f]
    for box in part_ii_boxes:
        row_num += 1
        _write_box_row(ws, row_num, box)

    # Part II Subtotal
    row_num += 1
    _write_box_row(ws, row_num, summary.long_term_total, is_subtotal=True)

    # Grand Total Row
    row_num += 2
    _write_box_row(ws, row_num, summary.grand_total, is_grand_total=True)

    # Auto-adjust column widths
    _autofit_columns(ws)


def _write_box_row(ws, row_num: int, box, is_subtotal: bool = False, is_grand_total: bool = False):
    """Writes a single row for a Form8949Box, applying red highlighting to wash sale amounts when present."""
    c_box = ws.cell(row=row_num, column=1, value=box.box_name)
    c_desc = ws.cell(row=row_num, column=2, value=box.description)
    c_proc = ws.cell(row=row_num, column=3, value=box.proceeds)
    c_cost = ws.cell(row=row_num, column=4, value=box.cost_basis)
    c_adj = ws.cell(row=row_num, column=5, value=box.adjustments)
    c_net = ws.cell(row=row_num, column=6, value=box.net_gain_loss)
    c_status = ws.cell(row=row_num, column=7, value="VALID" if box.is_parity_valid else "MISMATCH")

    font = BOLD_FONT if (is_subtotal or is_grand_total) else REGULAR_FONT
    border = DOUBLE_BOTTOM_BORDER if is_grand_total else THIN_BORDER

    for c in [c_box, c_desc, c_proc, c_cost, c_adj, c_net, c_status]:
        c.font = font
        c.border = border
        if is_subtotal:
            c.fill = SUBHEADER_FILL
        elif is_grand_total:
            c.fill = TOTAL_FILL

    # Number formats
    for c in [c_proc, c_cost, c_adj, c_net]:
        c.number_format = CURRENCY_FORMAT
        c.alignment = Alignment(horizontal="right", vertical="center")

    # Red highlighting for wash sale disallowances on Form 8949 Summary tab
    # Whenever wash sales are present, apply soft red fill (#FFC7CE) and bold dark red font (#9C0006)
    if box.wash_sale_disallowed > 0 or (box.adjustments > 0 and box.wash_sale_disallowed > 0):
        c_adj.fill = WASH_SALE_FILL
        c_adj.font = WASH_SALE_FONT

    c_box.alignment = Alignment(horizontal="center", vertical="center")
    c_status.alignment = Alignment(horizontal="center", vertical="center")
    c_status.font = SUCCESS_FONT if box.is_parity_valid else Font(name=FONT_FAMILY, size=10, bold=True, color="C00000")
    if box.is_parity_valid:
        c_status.fill = SUCCESS_FILL


def _build_breakdown_tab(ws, report: AggregationReport):
    """Builds the Per-Brokerage Breakdown tab."""
    ws.views.sheetView[0].showGridLines = True

    # Title block
    ws["A1"] = "Brokerage & Account Detail Breakdown"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Individual reporting figures and clearing affiliations for each ingested 1099 PDF statement"
    ws["A2"].font = Font(name=FONT_FAMILY, size=10, italic=True, color="595959")

    headers = [
        "Brokerage Brand",
        "Clearing Firm",
        "Account Number",
        "Source File",
        "ST Proceeds",
        "ST Cost Basis",
        "ST Adjustments",
        "ST Wash Sale (1g)",
        "ST Net Gain/(Loss)",
        "LT Proceeds",
        "LT Cost Basis",
        "LT Adjustments",
        "LT Wash Sale (1g)",
        "LT Net Gain/(Loss)",
        "Total Proceeds",
        "Total Cost Basis",
        "Total Adjustments",
        "Total Wash Sale (1g)",
        "Total Net Gain/(Loss)",
        "Parity Status",
    ]

    row_num = 4
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=row_num, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center" if col_idx in (3, 20) else "left", vertical="center")
        cell.border = THIN_BORDER
    ws.row_dimensions[row_num].height = 24

    for stmt in report.statements:
        row_num += 1
        s = stmt.form_8949
        st = s.short_term_total
        lt = s.long_term_total
        gt = s.grand_total
        st_wash = st.wash_sale_disallowed
        lt_wash = lt.wash_sale_disallowed
        tot_wash = s.total_wash_sale_disallowed

        cells = [
            ws.cell(row=row_num, column=1, value=stmt.brokerage_name),
            ws.cell(row=row_num, column=2, value=stmt.clearing_firm or "Direct / Self-Clearing"),
            ws.cell(row=row_num, column=3, value=stmt.account_number),
            ws.cell(row=row_num, column=4, value=stmt.source_file),
            ws.cell(row=row_num, column=5, value=st.proceeds),
            ws.cell(row=row_num, column=6, value=st.cost_basis),
            ws.cell(row=row_num, column=7, value=st.adjustments),
            ws.cell(row=row_num, column=8, value=st_wash),
            ws.cell(row=row_num, column=9, value=st.net_gain_loss),
            ws.cell(row=row_num, column=10, value=lt.proceeds),
            ws.cell(row=row_num, column=11, value=lt.cost_basis),
            ws.cell(row=row_num, column=12, value=lt.adjustments),
            ws.cell(row=row_num, column=13, value=lt_wash),
            ws.cell(row=row_num, column=14, value=lt.net_gain_loss),
            ws.cell(row=row_num, column=15, value=gt.proceeds),
            ws.cell(row=row_num, column=16, value=gt.cost_basis),
            ws.cell(row=row_num, column=17, value=gt.adjustments),
            ws.cell(row=row_num, column=18, value=tot_wash),
            ws.cell(row=row_num, column=19, value=gt.net_gain_loss),
            ws.cell(row=row_num, column=20, value="VALID" if stmt.parity_passed else "MISMATCH"),
        ]

        for idx, c in enumerate(cells, start=1):
            c.font = REGULAR_FONT
            c.border = THIN_BORDER
            if idx in range(5, 20):
                c.number_format = CURRENCY_FORMAT
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif idx in (3, 20):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

        # Red styling for wash sale columns when wash sales are present
        if st_wash > 0:
            cells[7].fill = WASH_SALE_FILL
            cells[7].font = WASH_SALE_FONT
        if lt_wash > 0:
            cells[12].fill = WASH_SALE_FILL
            cells[12].font = WASH_SALE_FONT
        if tot_wash > 0:
            cells[17].fill = WASH_SALE_FILL
            cells[17].font = WASH_SALE_FONT

        status_cell = cells[-1]
        status_cell.font = SUCCESS_FONT if stmt.parity_passed else Font(name=FONT_FAMILY, size=10, bold=True, color="C00000")
        if stmt.parity_passed:
            status_cell.fill = SUCCESS_FILL

    # Total Row across all brokerages
    row_num += 1
    agg = report.aggregate_8949
    agg_st = agg.short_term_total
    agg_lt = agg.long_term_total
    agg_gt = agg.grand_total
    agg_st_wash = agg_st.wash_sale_disallowed
    agg_lt_wash = agg_lt.wash_sale_disallowed
    agg_tot_wash = agg.total_wash_sale_disallowed

    total_cells = [
        ws.cell(row=row_num, column=1, value="Total All Accounts"),
        ws.cell(row=row_num, column=2, value=""),
        ws.cell(row=row_num, column=3, value=f"{len(report.statements)} accounts"),
        ws.cell(row=row_num, column=4, value=""),
        ws.cell(row=row_num, column=5, value=agg_st.proceeds),
        ws.cell(row=row_num, column=6, value=agg_st.cost_basis),
        ws.cell(row=row_num, column=7, value=agg_st.adjustments),
        ws.cell(row=row_num, column=8, value=agg_st_wash),
        ws.cell(row=row_num, column=9, value=agg_st.net_gain_loss),
        ws.cell(row=row_num, column=10, value=agg_lt.proceeds),
        ws.cell(row=row_num, column=11, value=agg_lt.cost_basis),
        ws.cell(row=row_num, column=12, value=agg_lt.adjustments),
        ws.cell(row=row_num, column=13, value=agg_lt_wash),
        ws.cell(row=row_num, column=14, value=agg_lt.net_gain_loss),
        ws.cell(row=row_num, column=15, value=agg_gt.proceeds),
        ws.cell(row=row_num, column=16, value=agg_gt.cost_basis),
        ws.cell(row=row_num, column=17, value=agg_gt.adjustments),
        ws.cell(row=row_num, column=18, value=agg_tot_wash),
        ws.cell(row=row_num, column=19, value=agg_gt.net_gain_loss),
        ws.cell(row=row_num, column=20, value="VALID" if report.all_parity_passed else "MISMATCH"),
    ]

    for idx, c in enumerate(total_cells, start=1):
        c.font = BOLD_FONT
        c.fill = TOTAL_FILL
        c.border = DOUBLE_BOTTOM_BORDER
        if idx in range(5, 20):
            c.number_format = CURRENCY_FORMAT
            c.alignment = Alignment(horizontal="right", vertical="center")
        elif idx in (1, 3, 20):
            c.alignment = Alignment(horizontal="center", vertical="center")

    # If aggregate has wash sales, also style total cells
    if agg_st_wash > 0:
        total_cells[7].fill = WASH_SALE_FILL
        total_cells[7].font = WASH_SALE_FONT
    if agg_lt_wash > 0:
        total_cells[12].fill = WASH_SALE_FILL
        total_cells[12].font = WASH_SALE_FONT
    if agg_tot_wash > 0:
        total_cells[17].fill = WASH_SALE_FILL
        total_cells[17].font = WASH_SALE_FONT

    _autofit_columns(ws)


def _build_int_div_tab(ws, report: AggregationReport):
    """Builds the Form 1099-INT and Form 1099-DIV Interest & Dividend Schedule tab."""
    ws.views.sheetView[0].showGridLines = True

    # Title block
    ws["A1"] = "Consolidated Form 1099-INT & Form 1099-DIV Schedule"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Breakdown of taxable interest, NY IT-201 S-102 exempt Treasury obligations, and qualified/ordinary dividends across accounts"
    ws["A2"].font = Font(name=FONT_FAMILY, size=10, italic=True, color="595959")

    headers = [
        "Brokerage Brand",
        "Clearing Firm",
        "Account Number",
        "Source File",
        "1099-INT Box 1 (Taxable Interest)",
        "1099-INT Box 3 (U.S. Treasury / NY IT-201 S-102)",
        "1099-INT Box 8 (Tax-Exempt Interest)",
        "Total 1099-INT Interest",
        "1099-DIV Box 1a (Total Ordinary Dividends)",
        "1099-DIV Box 1b (Qualified Dividends)",
        "1099-DIV Box 5 (Section 199A Dividends)",
    ]

    row_num = 4
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=row_num, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center" if col_idx == 3 else "left", vertical="center")
        cell.border = THIN_BORDER
    ws.row_dimensions[row_num].height = 24

    for stmt in report.statements:
        row_num += 1
        s_int = stmt.form_1099_int
        s_div = stmt.form_1099_div

        cells = [
            ws.cell(row=row_num, column=1, value=stmt.brokerage_name),
            ws.cell(row=row_num, column=2, value=stmt.clearing_firm or "Direct / Self-Clearing"),
            ws.cell(row=row_num, column=3, value=stmt.account_number),
            ws.cell(row=row_num, column=4, value=stmt.source_file),
            ws.cell(row=row_num, column=5, value=s_int.box_1_interest),
            ws.cell(row=row_num, column=6, value=s_int.box_3_us_treasury),
            ws.cell(row=row_num, column=7, value=s_int.box_8_tax_exempt),
            ws.cell(row=row_num, column=8, value=s_int.total_interest),
            ws.cell(row=row_num, column=9, value=s_div.box_1a_ordinary_dividends),
            ws.cell(row=row_num, column=10, value=s_div.box_1b_qualified_dividends),
            ws.cell(row=row_num, column=11, value=s_div.box_5_section_199a),
        ]

        for idx, c in enumerate(cells, start=1):
            c.font = REGULAR_FONT
            c.border = THIN_BORDER
            if idx in range(5, 12):
                c.number_format = CURRENCY_FORMAT
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif idx == 3:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

    # Total Row across all accounts
    row_num += 1
    agg_int = report.aggregate_int
    agg_div = report.aggregate_div

    total_cells = [
        ws.cell(row=row_num, column=1, value="Total All Accounts"),
        ws.cell(row=row_num, column=2, value=""),
        ws.cell(row=row_num, column=3, value=f"{len(report.statements)} accounts"),
        ws.cell(row=row_num, column=4, value=""),
        ws.cell(row=row_num, column=5, value=agg_int.box_1_interest),
        ws.cell(row=row_num, column=6, value=agg_int.box_3_us_treasury),
        ws.cell(row=row_num, column=7, value=agg_int.box_8_tax_exempt),
        ws.cell(row=row_num, column=8, value=agg_int.total_interest),
        ws.cell(row=row_num, column=9, value=agg_div.box_1a_ordinary_dividends),
        ws.cell(row=row_num, column=10, value=agg_div.box_1b_qualified_dividends),
        ws.cell(row=row_num, column=11, value=agg_div.box_5_section_199a),
    ]

    for idx, c in enumerate(total_cells, start=1):
        c.font = BOLD_FONT
        c.fill = TOTAL_FILL
        c.border = DOUBLE_BOTTOM_BORDER
        if idx in range(5, 12):
            c.number_format = CURRENCY_FORMAT
            c.alignment = Alignment(horizontal="right", vertical="center")
        elif idx in (1, 3):
            c.alignment = Alignment(horizontal="center", vertical="center")

    # Informational notes block
    row_num += 2
    cell_note_hdr = ws.cell(row=row_num, column=1, value="Key Tax Preparation & Filing Notes:")
    cell_note_hdr.font = Font(name=FONT_FAMILY, size=10, bold=True, color="1F4E79")

    notes = [
        "1. NY State Income Tax Subtraction (Form IT-201 Line 28, Code S-102): Form 1099-INT Box 3 (Interest on U.S. Savings Bonds and Treasury Obligations) is taxable on your federal return but is legally exempt from New York State and local income taxes pursuant to 31 U.S.C. § 3124. Deduct this total on Form IT-201.",
        "2. Federal Schedule B (Form 1040): Form 1099-INT Box 1 (Taxable Interest) is reported on Form 1040 Line 2b (and Schedule B Part I if taxable interest exceeds $1,500). Form 1099-INT Box 8 is reported on Line 2a as tax-exempt interest.",
        "3. Federal Qualified Dividends (Form 1040 Line 3a): Form 1099-DIV Box 1b (Qualified Dividends) is taxed at preferential long-term capital gains rates (0%, 15%, or 20%). Form 1099-DIV Box 1a is reported on Line 3b.",
        "4. Section 199A Dividends (Box 5): Eligible for the 20% Qualified Business Income (QBI) deduction on Form 8995 or Form 8995-A.",
    ]

    for n in notes:
        row_num += 1
        c_n = ws.cell(row=row_num, column=1, value=n)
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=len(headers))
        c_n.font = Font(name=FONT_FAMILY, size=9, italic=True, color="404040")

    _autofit_columns(ws)


def _autofit_columns(ws):
    """Sets column widths based on maximum contents."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if cell.number_format == CURRENCY_FORMAT and isinstance(cell.value, (int, float)):
                val_str = f"${cell.value:,.2f}"
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
