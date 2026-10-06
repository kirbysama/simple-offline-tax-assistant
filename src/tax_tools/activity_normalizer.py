"""Activity Normalizer for multi-bank credit card and banking statements."""

import csv
import io
import os
import re
from pathlib import Path
from typing import BinaryIO, List, Optional, Tuple, Union

import openpyxl

from tax_tools.sales_tax_models import BankDialect, NormalizedTransaction, TransactionTreatment


class ActivityNormalizer:
    """Auto-detects bank schemas and normalizes CSV/Excel transactions into standard data structures."""

    def detect_dialect(self, file_path_or_name: Union[str, Path], header_row: Optional[List[str]] = None) -> Tuple[BankDialect, str, str]:
        """Detect bank dialect, institution name, and account identifier.
        
        Returns:
            Tuple of (BankDialect, source_bank, source_account)
        """
        filename = Path(file_path_or_name).name
        fn_lower = filename.lower()

        if header_row is None:
            header_row = self._peek_header(file_path_or_name)

        headers_normalized = [h.strip().lower() for h in header_row if h]
        h_str = " ".join(headers_normalized)

        # Bank of America
        if "reference number" in h_str and "payee" in h_str:
            return BankDialect.BOFA, "Bank of America", self._extract_account_from_filename(filename, "Bank of America")

        # Apple Card
        if "clearing date" in h_str and ("merchant" in h_str or "amount (usd)" in h_str):
            return BankDialect.APPLE_CARD, "Apple Card", "Apple Card"

        # Capital One
        if "card no." in h_str or "card number" in h_str:
            return BankDialect.CAPITAL_ONE, "Capital One", self._extract_account_from_filename(filename, "Capital One")

        # Discover
        if "trans. date" in h_str or ("trans date" in h_str and "category" in h_str and "amount" in h_str):
            return BankDialect.DISCOVER, "Discover", self._extract_account_from_filename(filename, "Discover")

        # Chase Credit vs Chase Checking
        if "transaction date" in h_str and "post date" in h_str and "type" in h_str:
            return BankDialect.CHASE_CREDIT, "Chase", self._extract_account_from_filename(filename, "Chase")

        if "posting date" in h_str and "details" in h_str and "balance" in h_str:
            return BankDialect.CHASE_CHECKING, "Chase", self._extract_account_from_filename(filename, "Chase")

        # Amex
        if "date" in headers_normalized and (
            "extended details" in h_str or "appears on your statement as" in h_str or "reference" in h_str
        ):
            return BankDialect.AMEX, "American Express", self._extract_account_from_filename(filename, "American Express")

        # Citi
        if "debit" in headers_normalized and "credit" in headers_normalized:
            return BankDialect.CITI, "Citi", self._extract_account_from_filename(filename, "Citi")

        # Fallbacks using filename keywords
        if "chase" in fn_lower:
            if "check" in fn_lower:
                return BankDialect.CHASE_CHECKING, "Chase", self._extract_account_from_filename(filename, "Chase")
            return BankDialect.CHASE_CREDIT, "Chase", self._extract_account_from_filename(filename, "Chase")
        if "amex" in fn_lower or "american express" in fn_lower:
            return BankDialect.AMEX, "American Express", self._extract_account_from_filename(filename, "American Express")
        if "citi" in fn_lower:
            return BankDialect.CITI, "Citi", self._extract_account_from_filename(filename, "Citi")
        if "discover" in fn_lower:
            return BankDialect.DISCOVER, "Discover", self._extract_account_from_filename(filename, "Discover")
        if "capital one" in fn_lower or "capitalone" in fn_lower:
            return BankDialect.CAPITAL_ONE, "Capital One", self._extract_account_from_filename(filename, "Capital One")

        return BankDialect.GENERIC, "Generic Bank", self._extract_account_from_filename(filename, "Account")

    def _extract_account_from_filename(self, filename: str, default_bank: str) -> str:
        """Extract clean account name and/or masked number from filename."""
        base = Path(filename).stem
        # Remove timestamps like _Activity20250101_20251231_20260317
        base_clean = re.sub(r"_Activity[0-9_]+", "", base, flags=re.IGNORECASE)
        base_clean = re.sub(r" Activity [0-9]+", "", base_clean, flags=re.IGNORECASE)
        base_clean = re.sub(r"-Transactions-[0-9]+", "", base_clean, flags=re.IGNORECASE)
        # Look for 4 digits
        m = re.search(r"(\d{4})", base_clean)
        digits = f" (...{m.group(1)})" if m else ""
        
        # Clean specific patterns
        if "Freedom" in base_clean:
            return f"Freedom Card{digits}"
        if "Sapphire" in base_clean:
            return f"Sapphire Preferred{digits}"
        if "Personal Checking" in base_clean or "Checking" in base_clean:
            return f"Personal Checking{digits}"
        if "Blue Cash" in base_clean:
            return f"Blue Cash Everyday{digits}"
        if "Double Cash" in base_clean:
            return f"Double Cash{digits}"
        if "Discover" in base_clean:
            return f"Discover Card{digits}"

        return base_clean.replace("_", " ").strip() or f"{default_bank} Account"

    def _peek_header(self, file_path_or_name: Union[str, Path]) -> List[str]:
        """Read the header row of a CSV or Excel file."""
        path = Path(file_path_or_name)
        if not path.exists():
            return []

        if path.suffix.lower() in (".xlsx", ".xls"):
            try:
                wb = openpyxl.load_workbook(path, read_only=True)
                ws = wb.active
                for row in ws.iter_rows(values_only=True):
                    if any(row):
                        return [str(c or "").strip() for c in row]
            except Exception:
                return []
            return []

        # Read CSV with various encodings
        for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            try:
                with open(path, "r", encoding=enc, errors="replace") as f:
                    reader = csv.reader(f)
                    for row in reader:
                        if any(c.strip() for c in row):
                            return [c.strip() for c in row]
            except Exception:
                continue
        return []

    def parse_file(self, file_input: Union[str, Path, BinaryIO], filename: Optional[str] = None) -> List[NormalizedTransaction]:
        """Parse an activity CSV or Excel file into normalized transactions."""
        if isinstance(file_input, (str, Path)):
            file_path = Path(file_input)
            fname = file_path.name
            suffix = file_path.suffix.lower()
            if suffix in (".xlsx", ".xls"):
                return self._parse_excel(file_path, fname)
            else:
                return self._parse_csv_file(file_path, fname)
        else:
            fname = filename or "activity.csv"
            if fname.lower().endswith((".xlsx", ".xls")):
                return self._parse_excel_stream(file_input, fname)
            else:
                return self._parse_csv_stream(file_input, fname)

    def _parse_excel(self, path: Path, filename: str) -> List[NormalizedTransaction]:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(c or "").strip() for c in rows[0]]
        str_rows = [[str(c if c is not None else "").strip() for c in r] for r in rows[1:]]
        return self._normalize_rows(header, str_rows, filename)

    def _parse_excel_stream(self, stream: BinaryIO, filename: str) -> List[NormalizedTransaction]:
        wb = openpyxl.load_workbook(stream, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(c or "").strip() for c in rows[0]]
        str_rows = [[str(c if c is not None else "").strip() for c in r] for r in rows[1:]]
        return self._normalize_rows(header, str_rows, filename)

    def _parse_csv_file(self, path: Path, filename: str) -> List[NormalizedTransaction]:
        for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            try:
                with open(path, "r", encoding=enc, errors="replace") as f:
                    reader = csv.reader(f)
                    rows = list(reader)
                    if rows:
                        header = [c.strip() for c in rows[0]]
                        return self._normalize_rows(header, rows[1:], filename)
            except Exception:
                continue
        return []

    def _parse_csv_stream(self, stream: BinaryIO, filename: str) -> List[NormalizedTransaction]:
        content = stream.read()
        if isinstance(content, bytes):
            for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
                try:
                    text = content.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                text = content.decode("utf-8", errors="replace")
        else:
            text = content

        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        if not rows:
            return []
        header = [c.strip() for c in rows[0]]
        return self._normalize_rows(header, rows[1:], filename)

    def _normalize_rows(self, header: List[str], rows: List[List[str]], filename: str) -> List[NormalizedTransaction]:
        """Convert raw tabular rows into normalized transactions based on dialect."""
        dialect, bank, account = self.detect_dialect(filename, header)
        normalized_headers = {h.strip().lower(): idx for idx, h in enumerate(header)}
        
        txs: List[NormalizedTransaction] = []

        for row in rows:
            if not row or not any(c.strip() for c in row):
                continue
            
            # Helper to get column by name
            def get_col(*names: str) -> str:
                for name in names:
                    n_clean = name.lower()
                    if n_clean in normalized_headers and normalized_headers[n_clean] < len(row):
                        return row[normalized_headers[n_clean]].strip()
                return ""

            try:
                t = self._map_row(dialect, bank, account, row, get_col)
                if t:
                    txs.append(t)
            except Exception:
                continue

        return txs

    def _map_row(self, dialect: BankDialect, bank: str, account: str, row: List[str], get_col) -> Optional[NormalizedTransaction]:
        """Map a single row according to bank dialect."""
        def parse_float(val: str) -> float:
            if not val:
                return 0.0
            clean = val.replace("$", "").replace(",", "").strip()
            return float(clean) if clean else 0.0

        if dialect == BankDialect.CHASE_CREDIT:
            date = get_col("transaction date", "post date")
            desc = get_col("description")
            cat = get_col("category")
            tx_type = get_col("type")
            memo = get_col("memo")
            raw_amt = parse_float(get_col("amount"))
            
            # Chase convention: Sale is negative (-98.89), Payment is positive (19.99), Return is positive (183.74)
            is_refund = False
            if tx_type.lower() == "return":
                is_refund = True
                norm_amt = -abs(raw_amt)
            elif tx_type.lower() == "payment":
                norm_amt = abs(raw_amt)
            elif tx_type.lower() == "fee":
                norm_amt = abs(raw_amt)
            else: # Sale
                norm_amt = abs(raw_amt)

            return NormalizedTransaction(
                date=date,
                description=desc,
                amount=norm_amt,
                raw_amount=raw_amt,
                source_bank=bank,
                source_account=account,
                category=cat or "General Retail",
                is_refund=is_refund,
                memo=memo,
            )

        elif dialect == BankDialect.CHASE_CHECKING:
            date = get_col("posting date", "date")
            desc = get_col("description")
            tx_type = get_col("type")
            details = get_col("details")
            raw_amt = parse_float(get_col("amount"))
            
            # Debits are negative in checking (-100.00)
            norm_amt = abs(raw_amt)
            is_refund = (details.upper() == "CREDIT" and tx_type not in ("ACH_CREDIT", "DEPOSIT", "CHECK_DEPOSIT"))

            return NormalizedTransaction(
                date=date,
                description=desc,
                amount=norm_amt if not is_refund else -norm_amt,
                raw_amount=raw_amt,
                source_bank=bank,
                source_account=account,
                category=tx_type or "Banking",
                is_refund=is_refund,
            )

        elif dialect == BankDialect.AMEX:
            date = get_col("date")
            desc = get_col("description")
            cat = get_col("category")
            raw_amt = parse_float(get_col("amount"))
            
            # Amex convention: charges are positive (11.49), payments & returns are negative (-14.10, -15.00)
            desc_upper = desc.upper()
            is_payment = "PAYMENT" in desc_upper or "AUTOPAY" in desc_upper
            is_refund = (raw_amt < 0) and not is_payment

            return NormalizedTransaction(
                date=date,
                description=desc,
                amount=raw_amt if is_refund else abs(raw_amt),
                raw_amount=raw_amt,
                source_bank=bank,
                source_account=account,
                category=cat or "General Retail",
                is_refund=is_refund,
            )

        elif dialect == BankDialect.CITI:
            date = get_col("date")
            desc = get_col("description")
            debit_str = get_col("debit")
            credit_str = get_col("credit")
            
            debit_val = parse_float(debit_str)
            credit_val = parse_float(credit_str)

            desc_upper = desc.upper()
            is_payment = "AUTOPAY" in desc_upper or "PAYMENT" in desc_upper

            if debit_val > 0:
                raw_amt = debit_val
                norm_amt = debit_val
                is_refund = False
            else:
                raw_amt = credit_val
                norm_amt = -abs(credit_val) if not is_payment else abs(credit_val)
                is_refund = not is_payment

            return NormalizedTransaction(
                date=date,
                description=desc,
                amount=norm_amt,
                raw_amount=raw_amt,
                source_bank=bank,
                source_account=account,
                category="General Retail",
                is_refund=is_refund,
            )

        elif dialect == BankDialect.DISCOVER:
            date = get_col("trans. date", "trans date", "date")
            desc = get_col("description")
            cat = get_col("category")
            raw_amt = parse_float(get_col("amount"))

            desc_upper = desc.upper()
            is_payment = "DIRECTPAY" in desc_upper or "PAYMENT" in desc_upper or (cat and "Payments and Credits" in cat)
            is_refund = (raw_amt < 0) and not is_payment

            return NormalizedTransaction(
                date=date,
                description=desc,
                amount=raw_amt if is_refund else abs(raw_amt),
                raw_amount=raw_amt,
                source_bank=bank,
                source_account=account,
                category=cat or "General Retail",
                is_refund=is_refund,
            )

        elif dialect == BankDialect.CAPITAL_ONE:
            date = get_col("transaction date", "posted date", "date")
            desc = get_col("description")
            cat = get_col("category")
            card_no = get_col("card no.", "card number")
            if card_no:
                account = f"Capital One (...{card_no[-4:]})"

            debit_str = get_col("debit")
            credit_str = get_col("credit")
            amt_str = get_col("amount")

            if debit_str or credit_str:
                debit_val = parse_float(debit_str)
                credit_val = parse_float(credit_str)
                if debit_val > 0:
                    raw_amt = debit_val
                    norm_amt = debit_val
                    is_refund = False
                else:
                    raw_amt = credit_val
                    is_payment = "PAYMENT" in desc.upper() or "AUTOPAY" in desc.upper()
                    norm_amt = -abs(credit_val) if not is_payment else abs(credit_val)
                    is_refund = not is_payment
            else:
                raw_amt = parse_float(amt_str)
                is_payment = "PAYMENT" in desc.upper() or "AUTOPAY" in desc.upper()
                is_refund = (raw_amt < 0) and not is_payment
                norm_amt = raw_amt if is_refund else abs(raw_amt)

            return NormalizedTransaction(
                date=date,
                description=desc,
                amount=norm_amt,
                raw_amount=raw_amt,
                source_bank=bank,
                source_account=account,
                category=cat or "General Retail",
                is_refund=is_refund,
            )

        elif dialect == BankDialect.APPLE_CARD:
            date = get_col("transaction date", "date")
            desc = get_col("description", "merchant")
            cat = get_col("category")
            raw_amt = parse_float(get_col("amount (usd)", "amount"))
            tx_type = get_col("type")

            is_payment = tx_type.lower() == "payment" or "PAYMENT" in desc.upper()
            is_refund = (raw_amt < 0) and not is_payment

            return NormalizedTransaction(
                date=date,
                description=desc,
                amount=raw_amt if is_refund else abs(raw_amt),
                raw_amount=raw_amt,
                source_bank=bank,
                source_account=account,
                category=cat or "General Retail",
                is_refund=is_refund,
            )

        else: # GENERIC
            date = get_col("date", "transaction date", "trans date", "posting date")
            desc = get_col("description", "payee", "merchant", "details")
            cat = get_col("category")
            
            debit_str = get_col("debit")
            credit_str = get_col("credit")
            amt_str = get_col("amount")

            if debit_str or credit_str:
                debit_val = parse_float(debit_str)
                credit_val = parse_float(credit_str)
                if debit_val > 0:
                    raw_amt = debit_val
                    norm_amt = debit_val
                    is_refund = False
                else:
                    raw_amt = credit_val
                    is_payment = "PAYMENT" in desc.upper() or "AUTOPAY" in desc.upper()
                    norm_amt = -abs(credit_val) if not is_payment else abs(credit_val)
                    is_refund = not is_payment
            else:
                raw_amt = parse_float(amt_str)
                is_payment = "PAYMENT" in desc.upper() or "AUTOPAY" in desc.upper()
                is_refund = (raw_amt < 0) and not is_payment
                norm_amt = raw_amt if is_refund else abs(raw_amt)

            return NormalizedTransaction(
                date=date,
                description=desc,
                amount=norm_amt,
                raw_amount=raw_amt,
                source_bank=bank,
                source_account=account,
                category=cat or "General Retail",
                is_refund=is_refund,
            )
