"""Tests for Excel report generation and formatting."""

import os
import openpyxl
import pytest
from tax_tools.models import Form8949Box, Form8949Summary, BrokerageAccountStatement, AggregationReport
from tax_tools.excel_exporter import export_1099_excel


@pytest.fixture
def mock_report():
    from tax_tools.models import Form1099INT, Form1099DIV
    s1 = Form8949Summary(
        box_a=Form8949Box("Box A", "Short-term covered", proceeds=500.0, cost_basis=300.0, adjustments=0.0, net_gain_loss=200.0),
        box_d=Form8949Box("Box D", "Long-term covered", proceeds=1000.0, cost_basis=600.0, adjustments=0.0, net_gain_loss=400.0),
    )
    stmt1 = BrokerageAccountStatement(
        brokerage_name="Mock Brokerage",
        account_number="***-1234",
        clearing_firm="Mock Clearing",
        source_file="mock.pdf",
        form_8949=s1,
        form_1099_int=Form1099INT(box_1_interest=50.0, box_3_us_treasury=25.0, box_8_tax_exempt=10.0),
        form_1099_div=Form1099DIV(box_1a_ordinary_dividends=120.0, box_1b_qualified_dividends=80.0, box_5_section_199a=15.0),
    )
    return AggregationReport(statements=[stmt1])


def test_excel_export_structure(mock_report, tmp_path):
    output_file = str(tmp_path / "test_1099_Summary.xlsx")
    exported_path = export_1099_excel(mock_report, output_path=output_file)
    assert os.path.exists(exported_path)

    wb = openpyxl.load_workbook(exported_path)
    sheet_names = wb.sheetnames
    assert "Form 8949 Summary" in sheet_names
    assert "Brokerage Breakdown" in sheet_names
    assert "Interest & Dividend Schedule" in sheet_names

    # Check Summary tab
    ws_summary = wb["Form 8949 Summary"]
    assert "Consolidated IRS Form 8949 / Schedule D Summary" in str(ws_summary["A1"].value)
    
    # Check that Box A values appear
    found_box_a = False
    for row in ws_summary.iter_rows(values_only=True):
        if row and row[0] == "Box A":
            found_box_a = True
            assert row[2] == 500.0   # Proceeds
            assert row[3] == 300.0   # Basis
            assert row[5] == 200.0   # Net
            assert row[6] == "VALID" # Parity
    assert found_box_a is True

    # Check Breakdown tab
    ws_breakdown = wb["Brokerage Breakdown"]
    assert "Brokerage & Account Detail Breakdown" in str(ws_breakdown["A1"].value)
    found_stmt = False
    for row in ws_breakdown.iter_rows(values_only=True):
        if row and row[0] == "Mock Brokerage":
            found_stmt = True
            assert row[1] == "Mock Clearing"
            assert row[2] == "***-1234"
            # ST Proceeds, Basis, Adjustments, Wash, Net
            assert row[4] == 500.0   # ST Proceeds
            assert row[5] == 300.0   # ST Basis
            assert row[6] == 0.0     # ST Adjustments
            assert row[7] == 0.0     # ST Wash Sale
            assert row[8] == 200.0   # ST Net
            # LT Proceeds, Basis, Adjustments, Wash, Net
            assert row[9] == 1000.0  # LT Proceeds
            assert row[10] == 600.0  # LT Basis
            assert row[11] == 0.0    # LT Adjustments
            assert row[12] == 0.0    # LT Wash Sale
            assert row[13] == 400.0  # LT Net
            # Totals
            assert row[14] == 1500.0 # Total Proceeds
            assert row[15] == 900.0  # Total Cost Basis
            assert row[16] == 0.0    # Total Adjustments
            assert row[17] == 0.0    # Total Wash Sale (1g)
            assert row[18] == 600.0  # Total Net Gain/(Loss)
            assert row[19] == "VALID" # Parity Status
    assert found_stmt is True


def test_excel_wash_sale_red_highlighting(tmp_path):
    """Verifies that soft red fill (#FFC7CE) and bold dark red font (#9C0006) are applied when wash sales exist."""
    from tax_tools.excel_exporter import WASH_SALE_FILL, WASH_SALE_FONT
    s1 = Form8949Summary(
        box_a=Form8949Box("Box A", "Short-term covered", proceeds=15000.0, cost_basis=30000.0, adjustments=7500.0, net_gain_loss=-7500.0, wash_sale_disallowed=7500.0),
        box_b=Form8949Box("Box B", "Short-term non-covered", proceeds=500.0, cost_basis=400.0, adjustments=0.0, net_gain_loss=100.0, wash_sale_disallowed=0.0),
    )
    s2 = Form8949Summary(
        box_a=Form8949Box("Box A", "Clean account", proceeds=1000.0, cost_basis=800.0, adjustments=0.0, net_gain_loss=200.0, wash_sale_disallowed=0.0),
    )
    stmt1 = BrokerageAccountStatement(
        brokerage_name="Robinhood",
        account_number="90000200",
        clearing_firm="Robinhood Securities LLC",
        source_file="rh.pdf",
        form_8949=s1,
    )
    stmt2 = BrokerageAccountStatement(
        brokerage_name="Clean Brokerage",
        account_number="CLN-001",
        clearing_firm="Direct",
        source_file="clean.pdf",
        form_8949=s2,
    )
    report = AggregationReport(statements=[stmt1, stmt2])

    output_file = str(tmp_path / "test_wash_highlight.xlsx")
    export_1099_excel(report, output_path=output_file)

    wb = openpyxl.load_workbook(output_file)

    # 1. Inspect Form 8949 Summary tab
    ws_summary = wb["Form 8949 Summary"]
    box_a_row_idx = None
    box_b_row_idx = None
    for r in range(1, ws_summary.max_row + 1):
        val = ws_summary.cell(row=r, column=1).value
        if val == "Box A":
            box_a_row_idx = r
        elif val == "Box B":
            box_b_row_idx = r

    assert box_a_row_idx is not None
    assert box_b_row_idx is not None

    # Box A has wash sales -> column 5 (Adjustments) should have soft red fill and dark red font
    box_a_adj_cell = ws_summary.cell(row=box_a_row_idx, column=5)
    assert box_a_adj_cell.fill.start_color.rgb == "00FFC7CE" or box_a_adj_cell.fill.start_color.rgb == "FFC7CE"
    assert box_a_adj_cell.font.color.rgb == "009C0006" or box_a_adj_cell.font.color.rgb == "9C0006"
    assert box_a_adj_cell.font.bold is True

    # Box B has zero wash sales -> column 5 should NOT be red highlighted
    box_b_adj_cell = ws_summary.cell(row=box_b_row_idx, column=5)
    assert box_b_adj_cell.fill != WASH_SALE_FILL
    assert box_b_adj_cell.font.color is None or box_b_adj_cell.font.color.rgb != "009C0006"

    # 2. Inspect Brokerage Breakdown tab
    ws_breakdown = wb["Brokerage Breakdown"]
    rh_row_idx = None
    clean_row_idx = None
    for r in range(1, ws_breakdown.max_row + 1):
        val = ws_breakdown.cell(row=r, column=1).value
        if val == "Robinhood":
            rh_row_idx = r
        elif val == "Clean Brokerage":
            clean_row_idx = r

    assert rh_row_idx is not None
    assert clean_row_idx is not None

    # Robinhood has ST wash sales -> column 8 (ST Wash Sale) & column 18 (Total Wash Sale) should have soft red fill & dark red bold font
    rh_st_wash_cell = ws_breakdown.cell(row=rh_row_idx, column=8)
    assert rh_st_wash_cell.value == 7500.0
    assert rh_st_wash_cell.fill.start_color.rgb == "00FFC7CE" or rh_st_wash_cell.fill.start_color.rgb == "FFC7CE"
    assert rh_st_wash_cell.font.color.rgb == "009C0006" or rh_st_wash_cell.font.color.rgb == "9C0006"
    assert rh_st_wash_cell.font.bold is True

    rh_wash_cell = ws_breakdown.cell(row=rh_row_idx, column=18)
    assert rh_wash_cell.value == 7500.0
    assert rh_wash_cell.fill.start_color.rgb == "00FFC7CE" or rh_wash_cell.fill.start_color.rgb == "FFC7CE"
    assert rh_wash_cell.font.color.rgb == "009C0006" or rh_wash_cell.font.color.rgb == "9C0006"
    assert rh_wash_cell.font.bold is True

    # Clean Brokerage has 0 wash sales -> columns 8 and 18 should NOT be red highlighted
    clean_st_wash_cell = ws_breakdown.cell(row=clean_row_idx, column=8)
    assert clean_st_wash_cell.value == 0.0
    assert clean_st_wash_cell.fill != WASH_SALE_FILL
    assert clean_st_wash_cell.font.color is None or clean_st_wash_cell.font.color.rgb != "009C0006"

    clean_wash_cell = ws_breakdown.cell(row=clean_row_idx, column=18)
    assert clean_wash_cell.value == 0.0
    assert clean_wash_cell.fill != WASH_SALE_FILL
    assert clean_wash_cell.font.color is None or clean_wash_cell.font.color.rgb != "009C0006"
